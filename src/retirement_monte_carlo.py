from pathlib import Path
import textwrap

script = r'''
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl
import scipy.stats as st
import matplotlib.pyplot as plt

# -----------------------------
# Configuration
# -----------------------------
BASE_DIR = Path(__file__).resolve().parents[1]
INPUT_XLSX = BASE_DIR / "data" / "retirement_planning_data_2026.xlsx"
OUTPUT_DIR = BASE_DIR / "outputs"
OUTPUT_SUMMARY_CSV = OUTPUT_DIR / "retirement_simulation_summary.csv"
OUTPUT_CHART_PNG = OUTPUT_DIR / "retirement_overlay_chart.png"

MONTHLY_CONTRIBUTION = 2000
N_MONTHS = 72
GOAL = 170000
PORTFOLIO_WEIGHTS = [0.00, 0.25, 0.50, 0.75, 1.00]
N_SIMS = 100000
SEED = 42

# Candidate distributions to fit to monthly return data
CANDIDATES = {
    "norm": st.norm,
    "t": st.t,
    "laplace": st.laplace,
    "logistic": st.logistic,
}


def load_returns_from_workbook(xlsx_path: Path):
    """
    Reads the historical monthly returns from the 'model' sheet.
    Based on your workbook:
    - S&P 500 monthly returns are in B11:B250
    - Treasury monthly returns are in C11:C250
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["model"]

    sp500 = np.array([ws[f"B{r}"].value for r in range(11, 251)], dtype=float)
    treasury = np.array([ws[f"C{r}"].value for r in range(11, 251)], dtype=float)

    return sp500, treasury


def fit_best_distribution(data: np.ndarray):
    """
    Fit several distributions and choose the one with the lowest AIC.
    Returns:
        best_fit: (name, distribution_object, fitted_params, aic)
        all_fits: list of all fits sorted by AIC
    """
    fits = []

    for name, dist in CANDIDATES.items():
        params = dist.fit(data)
        log_likelihood = np.sum(dist.logpdf(data, *params))
        aic = 2 * len(params) - 2 * log_likelihood
        fits.append((name, dist, params, aic))

    fits.sort(key=lambda x: x[3])
    return fits[0], fits


def simulate_final_balances(
    sp500_dist,
    sp500_params,
    treasury_dist,
    treasury_params,
    n_sims=N_SIMS,
    n_months=N_MONTHS,
    monthly_contribution=MONTHLY_CONTRIBUTION,
    weights=PORTFOLIO_WEIGHTS,
    seed=SEED,
):
    """
    Simulate final account balances after 72 months.

    Assumption used here:
    - Each month's contribution is invested at the beginning of the month
    - Then that month's return is applied

    This matches a typical monthly investment setup and is easy to explain.
    """
    rng = np.random.default_rng(seed)

    # Generate the same market paths for every strategy so comparisons are fair
    sp500_returns = sp500_dist.rvs(*sp500_params, size=(n_sims, n_months), random_state=rng)
    treasury_returns = treasury_dist.rvs(*treasury_params, size=(n_sims, n_months), random_state=rng)

    final_balances = {}

    for w in weights:
        sp_balance = np.zeros(n_sims)
        treasury_balance = np.zeros(n_sims)

        sp_contribution = monthly_contribution * w
        treasury_contribution = monthly_contribution * (1 - w)

        for month in range(n_months):
            sp_balance = (sp_balance + sp_contribution) * (1 + sp500_returns[:, month])
            treasury_balance = (treasury_balance + treasury_contribution) * (1 + treasury_returns[:, month])

        final_balances[w] = sp_balance + treasury_balance

    return final_balances


def summarize_results(final_balances: dict, goal=GOAL):
    rows = []

    for w, balances in final_balances.items():
        rows.append(
            {
                "fraction_in_sp500": w,
                "mean_final_value": balances.mean(),
                "median_final_value": np.median(balances),
                "p5_final_value": np.quantile(balances, 0.05),
                "p95_final_value": np.quantile(balances, 0.95),
                "probability_reaching_goal": np.mean(balances >= goal),
            }
        )

    return pd.DataFrame(rows).sort_values("fraction_in_sp500").reset_index(drop=True)


def plot_overlay_cdf(final_balances: dict, goal=GOAL, output_path=OUTPUT_CHART_PNG):
    plt.figure(figsize=(9, 6))

    for w, balances in sorted(final_balances.items()):
        x = np.sort(balances)
        y = np.arange(1, len(x) + 1) / len(x)
        plt.plot(x, y, label=f"{int(w * 100)}% S&P 500")

    plt.axvline(goal, linestyle="--", label=f"Goal = ${goal:,.0f}")
    plt.xlabel("Final account value after 6 years ($)")
    plt.ylabel("Cumulative probability")
    plt.title("Overlay chart: cumulative distributions of final wealth")
    plt.legend()
    plt.tight_layout()
    plt.savefig(output_path, dpi=200)
    plt.close()


def print_fit_table(fits, series_name):
    print(f"\nDistribution fits for {series_name} monthly returns:")
    for name, _, params, aic in fits:
        print(f"  {name:<10} AIC = {aic:,.2f} | params = {params}")


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    if not INPUT_XLSX.exists():
        raise FileNotFoundError(
            f"Could not find the workbook at:\n{INPUT_XLSX}\n\n"
            "Put your Excel file in the repo's data/ folder."
        )

    sp500, treasury = load_returns_from_workbook(INPUT_XLSX)

    best_sp, all_sp = fit_best_distribution(sp500)
    best_tr, all_tr = fit_best_distribution(treasury)

    print_fit_table(all_sp, "S&P 500")
    print_fit_table(all_tr, "Treasury")

    sp_name, sp_dist, sp_params, sp_aic = best_sp
    tr_name, tr_dist, tr_params, tr_aic = best_tr

    print("\nChosen distributions:")
    print(f"  S&P 500  -> {sp_name} (AIC = {sp_aic:,.2f})")
    print(f"  Treasury -> {tr_name} (AIC = {tr_aic:,.2f})")

    final_balances = simulate_final_balances(
        sp500_dist=sp_dist,
        sp500_params=sp_params,
        treasury_dist=tr_dist,
        treasury_params=tr_params,
    )

    summary = summarize_results(final_balances)
    summary.to_csv(OUTPUT_SUMMARY_CSV, index=False)

    plot_overlay_cdf(final_balances, goal=GOAL, output_path=OUTPUT_CHART_PNG)

    print("\nSimulation summary:")
    print(summary.to_string(index=False, float_format=lambda x: f"{x:,.4f}"))

    print(f"\nSaved summary table to: {OUTPUT_SUMMARY_CSV}")
    print(f"Saved overlay chart to:  {OUTPUT_CHART_PNG}")


if __name__ == "__main__":
    main()
'''

requirements = """numpy
pandas
openpyxl
scipy
matplotlib
"""

readme_snippet = """# Retirement Monte Carlo Simulation

This project runs a Monte Carlo simulation for a retirement planning assignment.

## Setup

```bash
python -m venv .venv
source .venv/bin/activate
pip install -r requirements.txt